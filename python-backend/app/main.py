from __future__ import annotations

import asyncio
import csv
import io
import json
import math
import secrets
import time
from contextlib import asynccontextmanager, suppress
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

from fastapi import Body, Depends, FastAPI, File, Form, Header, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse

from .config import settings
from .database import db, dumps, hash_password, loads, new_id, now_iso, verify_password
from .generation_tasks import GenerationControl, generation_tasks
from .rag import agent, document_to_api, get_document, ingest_text, seed_product_knowledge, structure_chunks, tokenize
from .retrieval import close_vector_store, index_chunks, remove_chunks


def vector_rows(chunk_ids: list[str]) -> list[dict[str, Any]]:
    if not chunk_ids:
        return []
    placeholders = ",".join("?" for _ in chunk_ids)
    with db.connect() as connection:
        rows = connection.execute(
            f"""SELECT c.id chunk_id, c.doc_id document_id, c.content, c.content_hash,
                       d.doc_name document_name, kb.name knowledge_base
                FROM chunks c JOIN documents d ON d.id=c.doc_id
                JOIN knowledge_bases kb ON kb.id=c.kb_id
                WHERE c.id IN ({placeholders}) AND c.enabled=1 AND d.enabled=1""",
            chunk_ids,
        ).fetchall()
    return [dict(row) for row in rows]


def ok(data: Any = None) -> dict[str, Any]:
    return {"code": "0", "message": "success", "data": data}


def fail(message: str, status_code: int = 400) -> None:
    raise HTTPException(status_code=status_code, detail=message)


def page(records: list[Any], total: int, current: int, size: int) -> dict[str, Any]:
    return {
        "records": records,
        "total": total,
        "size": size,
        "current": current,
        "pages": math.ceil(total / size) if size else 0,
    }


def row_dict(row: Any) -> dict[str, Any]:
    return dict(row) if row is not None else {}


def extract_token(authorization: str | None) -> str:
    if not authorization:
        return ""
    if authorization.lower().startswith("bearer "):
        return authorization[7:].strip()
    return authorization.strip()


def current_user(authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = extract_token(authorization)
    with db.connect() as connection:
        row = connection.execute(
            """SELECT u.* FROM users u JOIN tokens t ON t.user_id=u.id WHERE t.token=?""",
            (token,),
        ).fetchone()
    if not row:
        fail("未登录或登录已过期", 401)
    return dict(row)


def admin_user(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    if user["role"] != "admin":
        fail("需要管理员权限", 403)
    return user


def user_api(user: dict[str, Any], token: str | None = None) -> dict[str, Any]:
    result = {
        "userId": user["id"], "id": user["id"], "username": user["username"],
        "role": user["role"], "avatar": user.get("avatar"),
    }
    if token:
        result["token"] = token
    return result


def camel_user(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"], "username": row["username"], "role": row["role"],
        "avatar": row.get("avatar"), "createTime": row.get("create_time"), "updateTime": row.get("update_time"),
    }


@asynccontextmanager
async def lifespan(_: FastAPI):
    db.initialize()
    settings.upload_dir.mkdir(parents=True, exist_ok=True)
    seed_product_knowledge()
    try:
        yield
    finally:
        await generation_tasks.cancel_all()
        close_vector_store()


app = FastAPI(
    title="AI Product Assistant RAG - Python",
    description="Agentic RAG product assistant compatible with the existing React console.",
    version="1.0.0",
    lifespan=lifespan,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(HTTPException)
async def http_exception_handler(_: Any, exc: HTTPException) -> JSONResponse:
    return JSONResponse(
        status_code=exc.status_code,
        content={"code": str(exc.status_code), "message": str(exc.detail), "data": None},
    )


@app.get("/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "service": settings.app_name, "backend": "python-fastapi"}


PREFIX = settings.api_prefix


@app.post(PREFIX + "/auth/login")
def login(payload: dict[str, Any] = Body(...)) -> dict[str, Any]:
    username = str(payload.get("username", "")).strip()
    password = str(payload.get("password", ""))
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        if not row or not verify_password(password, row["password_hash"]):
            fail("用户名或密码错误", 401)
        token = secrets.token_urlsafe(32)
        connection.execute("INSERT INTO tokens VALUES (?, ?, ?)", (token, row["id"], now_iso()))
    return ok(user_api(dict(row), token))


@app.post(PREFIX + "/auth/register")
def register(payload: dict[str, Any] = Body(...)) -> dict[str, Any]:
    username = str(payload.get("username", "")).strip()
    password = str(payload.get("password", ""))
    if len(username) < 3 or len(password) < 6:
        fail("用户名至少 3 个字符，密码至少 6 个字符")
    timestamp = now_iso()
    try:
        with db.connect() as connection:
            user_id = new_id()
            connection.execute(
                "INSERT INTO users VALUES (?, ?, ?, 'user', NULL, ?, ?)",
                (user_id, username, hash_password(password), timestamp, timestamp),
            )
    except Exception as exc:
        if "UNIQUE" in str(exc):
            fail("用户名已存在")
        raise
    return ok("注册成功")


@app.post(PREFIX + "/auth/logout")
def logout(authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = extract_token(authorization)
    with db.connect() as connection:
        connection.execute("DELETE FROM tokens WHERE token=?", (token,))
    return ok()


@app.get(PREFIX + "/user/me")
def me(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    return ok(user_api(user))


@app.get(PREFIX + "/users")
def users(
    current: int = 1,
    size: int = 10,
    keyword: str | None = None,
    _: dict[str, Any] = Depends(admin_user),
) -> dict[str, Any]:
    where, values = "", []
    if keyword:
        where, values = " WHERE username LIKE ?", [f"%{keyword}%"]
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM users{where}", values).fetchone()[0]
        rows = connection.execute(
            f"SELECT * FROM users{where} ORDER BY create_time DESC LIMIT ? OFFSET ?",
            (*values, size, (current - 1) * size),
        ).fetchall()
    return ok(page([camel_user(dict(row)) for row in rows], total, current, size))


@app.post(PREFIX + "/users")
def create_user(payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    username = str(payload.get("username", "")).strip()
    password = str(payload.get("password", ""))
    if len(username) < 3 or len(password) < 6:
        fail("用户名至少 3 个字符，密码至少 6 个字符")
    timestamp, user_id = now_iso(), new_id()
    try:
        with db.connect() as connection:
            connection.execute(
                "INSERT INTO users VALUES (?, ?, ?, ?, ?, ?, ?)",
                (user_id, username, hash_password(password), payload.get("role", "user"), payload.get("avatar"), timestamp, timestamp),
            )
    except Exception as exc:
        if "UNIQUE" in str(exc):
            fail("用户名已存在")
        raise
    return ok(user_id)


@app.put(PREFIX + "/users/{user_id}")
def update_user(user_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    fields, values = [], []
    for key, column in (("username", "username"), ("role", "role"), ("avatar", "avatar")):
        if key in payload:
            fields.append(f"{column}=?")
            values.append(payload[key])
    if payload.get("password"):
        fields.append("password_hash=?")
        values.append(hash_password(payload["password"]))
    fields.append("update_time=?")
    values.extend([now_iso(), user_id])
    with db.connect() as connection:
        connection.execute(f"UPDATE users SET {', '.join(fields)} WHERE id=?", values)
    return ok()


@app.delete(PREFIX + "/users/{user_id}")
def delete_user(user_id: str, user: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    if user_id == user["id"]:
        fail("不能删除当前登录用户")
    with db.connect() as connection:
        connection.execute("DELETE FROM users WHERE id=?", (user_id,))
    return ok()


@app.put(PREFIX + "/user/password")
def change_password(payload: dict[str, Any] = Body(...), user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    if not verify_password(str(payload.get("currentPassword", "")), user["password_hash"]):
        fail("当前密码错误")
    password = str(payload.get("newPassword", ""))
    if len(password) < 6:
        fail("新密码至少 6 个字符")
    with db.connect() as connection:
        connection.execute("UPDATE users SET password_hash=?, update_time=? WHERE id=?", (hash_password(password), now_iso(), user["id"]))
    return ok()


def kb_api(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"], "name": row["name"], "embeddingModel": row["embedding_model"],
        "collectionName": row["collection_name"], "createdBy": row.get("created_by"),
        "documentCount": row.get("document_count", 0), "createTime": row.get("create_time"), "updateTime": row.get("update_time"),
    }


@app.get(PREFIX + "/knowledge-base/chunk-strategies")
def chunk_strategies(_: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return ok([
        {"value": "structure", "label": "结构感知切分", "defaultConfig": {"maxChars": 900, "overlap": 120}},
        {"value": "fixed", "label": "固定长度切分", "defaultConfig": {"maxChars": 700, "overlap": 100}},
    ])


@app.get(PREFIX + "/knowledge-base")
def knowledge_bases(
    current: int = 1, size: int = 10, name: str | None = None,
    _: dict[str, Any] = Depends(admin_user),
) -> dict[str, Any]:
    where, values = (" WHERE kb.name LIKE ?", [f"%{name}%"]) if name else ("", [])
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM knowledge_bases kb{where}", values).fetchone()[0]
        rows = connection.execute(
            f"""SELECT kb.*, (SELECT COUNT(*) FROM documents d WHERE d.kb_id=kb.id) document_count
                FROM knowledge_bases kb{where} ORDER BY kb.create_time DESC LIMIT ? OFFSET ?""",
            (*values, size, (current - 1) * size),
        ).fetchall()
    return ok(page([kb_api(dict(row)) for row in rows], total, current, size))


@app.post(PREFIX + "/knowledge-base")
def create_kb(payload: dict[str, Any] = Body(...), user: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    kb_id, timestamp = new_id(), now_iso()
    name = str(payload.get("name", "")).strip()
    if not name:
        fail("知识库名称不能为空")
    with db.connect() as connection:
        connection.execute(
            "INSERT INTO knowledge_bases VALUES (?, ?, ?, ?, ?, ?, ?)",
            (kb_id, name, payload.get("embeddingModel", "local-hybrid"), payload.get("collectionName", f"kb_{kb_id[:10]}"), user["username"], timestamp, timestamp),
        )
    return ok(kb_id)


@app.get(PREFIX + "/knowledge-base/docs/search")
def search_docs(keyword: str = "", limit: int = 8, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        rows = connection.execute(
            """SELECT d.id, d.kb_id, d.doc_name, kb.name kb_name FROM documents d
               JOIN knowledge_bases kb ON kb.id=d.kb_id WHERE d.doc_name LIKE ? LIMIT ?""",
            (f"%{keyword}%", limit),
        ).fetchall()
    return ok([{"id": row["id"], "kbId": row["kb_id"], "docName": row["doc_name"], "kbName": row["kb_name"]} for row in rows])


@app.get(PREFIX + "/knowledge-base/{kb_id}")
def knowledge_base(kb_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute(
            """SELECT kb.*, (SELECT COUNT(*) FROM documents d WHERE d.kb_id=kb.id) document_count
               FROM knowledge_bases kb WHERE kb.id=?""",
            (kb_id,),
        ).fetchone()
    if not row:
        fail("知识库不存在", 404)
    return ok(kb_api(dict(row)))


@app.put(PREFIX + "/knowledge-base/{kb_id}")
def update_kb(kb_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    fields, values = [], []
    for key, column in (("name", "name"), ("embeddingModel", "embedding_model")):
        if key in payload:
            fields.append(f"{column}=?")
            values.append(payload[key])
    if fields:
        values.extend([now_iso(), kb_id])
        with db.connect() as connection:
            connection.execute(f"UPDATE knowledge_bases SET {', '.join(fields)}, update_time=? WHERE id=?", values)
    return ok()


@app.delete(PREFIX + "/knowledge-base/{kb_id}")
def delete_kb(kb_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("DELETE FROM knowledge_bases WHERE id=?", (kb_id,))
    return ok()


@app.get(PREFIX + "/knowledge-base/{kb_id}/docs")
def documents(
    kb_id: str, current: int = 1, size: int = 10, status: str | None = None,
    keyword: str | None = None, _: dict[str, Any] = Depends(admin_user),
) -> dict[str, Any]:
    clauses, values = ["d.kb_id=?"], [kb_id]
    if status:
        clauses.append("d.status=?")
        values.append(status)
    if keyword:
        clauses.append("d.doc_name LIKE ?")
        values.append(f"%{keyword}%")
    where = " AND ".join(clauses)
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM documents d WHERE {where}", values).fetchone()[0]
        rows = connection.execute(
            f"""SELECT d.*, (SELECT COUNT(*) FROM chunks c WHERE c.doc_id=d.id) chunk_count
                FROM documents d WHERE {where} ORDER BY d.create_time DESC LIMIT ? OFFSET ?""",
            (*values, size, (current - 1) * size),
        ).fetchall()
    return ok(page([document_to_api(dict(row)) for row in rows], total, current, size))


async def upload_to_text(file: UploadFile) -> tuple[str, bytes]:
    content = await file.read()
    suffix = Path(file.filename or "document.txt").suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages), content
        except Exception as exc:
            fail(f"PDF 解析失败：{exc}")
    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(io.BytesIO(content))
            return "\n".join(paragraph.text for paragraph in document.paragraphs), content
        except Exception as exc:
            fail(f"Word 解析失败：{exc}")
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding), content
        except UnicodeDecodeError:
            continue
    fail("无法识别文件编码")


@app.post(PREFIX + "/knowledge-base/{kb_id}/docs/upload")
async def upload_document(
    kb_id: str,
    sourceType: str = Form("file"),
    file: UploadFile | None = File(default=None),
    sourceLocation: str | None = Form(default=None),
    chunkStrategy: str = Form("structure"),
    processMode: str = Form("chunk"),
    scheduleEnabled: bool = Form(False),
    scheduleCron: str | None = Form(default=None),
    pipelineId: str | None = Form(default=None),
    _: dict[str, Any] = Depends(admin_user),
) -> dict[str, Any]:
    with db.connect() as connection:
        if not connection.execute("SELECT 1 FROM knowledge_bases WHERE id=?", (kb_id,)).fetchone():
            fail("知识库不存在", 404)
    if sourceType == "url":
        if not sourceLocation:
            fail("URL 不能为空")
        try:
            import httpx

            async with httpx.AsyncClient(timeout=20, follow_redirects=True) as client:
                response = await client.get(sourceLocation)
                response.raise_for_status()
                text, raw = response.text, response.content
            name = Path(sourceLocation.split("?", 1)[0]).name or "URL 文档"
        except Exception as exc:
            fail(f"URL 获取失败：{exc}")
    else:
        if not file:
            fail("请选择文件")
        text, raw = await upload_to_text(file)
        name = file.filename or "document.txt"
        save_path = settings.upload_dir / f"{new_id()}_{Path(name).name}"
        save_path.write_bytes(raw)
        sourceLocation = str(save_path)
    result = ingest_text(kb_id, name, text, source_type=sourceType, source_location=sourceLocation, file_size=len(raw), chunk_strategy=chunkStrategy)
    with db.connect() as connection:
        connection.execute(
            "UPDATE documents SET process_mode=?, schedule_enabled=?, schedule_cron=?, pipeline_id=? WHERE id=?",
            (processMode, int(scheduleEnabled), scheduleCron, pipelineId, result["id"]),
        )
    return ok(get_document(result["id"]))


@app.get(PREFIX + "/knowledge-base/docs/{doc_id}")
def document(doc_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    try:
        return ok(get_document(doc_id))
    except KeyError:
        fail("文档不存在", 404)


@app.put(PREFIX + "/knowledge-base/docs/{doc_id}")
def update_document(doc_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    mapping = {
        "docName": "doc_name", "processMode": "process_mode", "chunkStrategy": "chunk_strategy",
        "chunkConfig": "chunk_config", "pipelineId": "pipeline_id", "sourceLocation": "source_location",
        "scheduleEnabled": "schedule_enabled", "scheduleCron": "schedule_cron",
    }
    fields, values = [], []
    for key, column in mapping.items():
        if key in payload:
            fields.append(f"{column}=?")
            values.append(payload[key])
    if fields:
        values.extend([now_iso(), doc_id])
        with db.connect() as connection:
            connection.execute(f"UPDATE documents SET {', '.join(fields)}, update_time=? WHERE id=?", values)
    return ok()


@app.post(PREFIX + "/knowledge-base/docs/{doc_id}/chunk")
@app.post(PREFIX + "/knowledge-base/docs/{doc_id}/chunks/rebuild")
def rebuild_document(doc_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
    if not row:
        fail("文档不存在", 404)
    ingest_text(row["kb_id"], row["doc_name"], row["raw_text"] or "", source_type=row["source_type"] or "file", source_location=row["source_location"], chunk_strategy=row["chunk_strategy"] or "structure", document_id=doc_id)
    return ok()


@app.patch(PREFIX + "/knowledge-base/docs/{doc_id}/enable")
def enable_document(doc_id: str, value: bool, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("UPDATE documents SET enabled=?, update_time=? WHERE id=?", (int(value), now_iso(), doc_id))
        chunk_ids = [row["id"] for row in connection.execute("SELECT id FROM chunks WHERE doc_id=?", (doc_id,))]
    if value:
        index_chunks(vector_rows(chunk_ids))
    else:
        remove_chunks(chunk_ids)
    return ok()


@app.delete(PREFIX + "/knowledge-base/docs/{doc_id}")
def delete_document(doc_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        chunk_ids = [row["id"] for row in connection.execute("SELECT id FROM chunks WHERE doc_id=?", (doc_id,))]
        connection.execute("DELETE FROM documents WHERE id=?", (doc_id,))
    remove_chunks(chunk_ids)
    return ok()


def chunk_api(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"], "kbId": row["kb_id"], "docId": row["doc_id"], "chunkIndex": row["chunk_index"],
        "content": row["content"], "contentHash": row["content_hash"], "charCount": row["char_count"],
        "tokenCount": row["token_count"], "enabled": row["enabled"], "createTime": row["create_time"], "updateTime": row["update_time"],
    }


@app.get(PREFIX + "/knowledge-base/docs/{doc_id}/chunks")
def chunks(doc_id: str, current: int = 1, size: int = 10, enabled: int | None = None, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    where, values = "doc_id=?", [doc_id]
    if enabled is not None:
        where += " AND enabled=?"
        values.append(enabled)
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM chunks WHERE {where}", values).fetchone()[0]
        rows = connection.execute(
            f"SELECT * FROM chunks WHERE {where} ORDER BY chunk_index LIMIT ? OFFSET ?",
            (*values, size, (current - 1) * size),
        ).fetchall()
    return ok(page([chunk_api(dict(row)) for row in rows], total, current, size))


@app.post(PREFIX + "/knowledge-base/docs/{doc_id}/chunks")
def create_chunk(doc_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    content = str(payload.get("content", "")).strip()
    if not content:
        fail("分块内容不能为空")
    with db.connect() as connection:
        document_row = connection.execute("SELECT kb_id FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not document_row:
            fail("文档不存在", 404)
        index = payload.get("index")
        if index is None:
            index = connection.execute("SELECT COALESCE(MAX(chunk_index), -1)+1 FROM chunks WHERE doc_id=?", (doc_id,)).fetchone()[0]
        chunk_id, timestamp = payload.get("chunkId") or new_id(), now_iso()
        digest = __import__("hashlib").sha256(content.encode()).hexdigest()
        connection.execute(
            "INSERT INTO chunks VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)",
            (chunk_id, document_row["kb_id"], doc_id, index, content, digest, len(content), len(tokenize(content)), timestamp, timestamp),
        )
        row = connection.execute("SELECT * FROM chunks WHERE id=?", (chunk_id,)).fetchone()
    index_chunks(vector_rows([chunk_id]))
    return ok(chunk_api(dict(row)))


@app.put(PREFIX + "/knowledge-base/docs/{doc_id}/chunks/{chunk_id}")
def update_chunk(doc_id: str, chunk_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    content = str(payload.get("content", "")).strip()
    digest = __import__("hashlib").sha256(content.encode()).hexdigest()
    with db.connect() as connection:
        connection.execute(
            "UPDATE chunks SET content=?, content_hash=?, char_count=?, token_count=?, update_time=? WHERE id=? AND doc_id=?",
            (content, digest, len(content), len(tokenize(content)), now_iso(), chunk_id, doc_id),
        )
    index_chunks(vector_rows([chunk_id]))
    return ok()


@app.delete(PREFIX + "/knowledge-base/docs/{doc_id}/chunks/{chunk_id}")
def delete_chunk(doc_id: str, chunk_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("DELETE FROM chunks WHERE id=? AND doc_id=?", (chunk_id, doc_id))
    remove_chunks([chunk_id])
    return ok()


@app.post(PREFIX + "/knowledge-base/docs/{doc_id}/chunks/{chunk_id}/enable")
def enable_chunk(doc_id: str, chunk_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return _set_chunk_state(doc_id, [chunk_id], 1)


@app.post(PREFIX + "/knowledge-base/docs/{doc_id}/chunks/{chunk_id}/disable")
def disable_chunk(doc_id: str, chunk_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return _set_chunk_state(doc_id, [chunk_id], 0)


def _set_chunk_state(doc_id: str, chunk_ids: list[str] | None, state: int) -> dict[str, Any]:
    with db.connect() as connection:
        if chunk_ids:
            placeholders = ",".join("?" for _ in chunk_ids)
            connection.execute(
                f"UPDATE chunks SET enabled=?, update_time=? WHERE doc_id=? AND id IN ({placeholders})",
                (state, now_iso(), doc_id, *chunk_ids),
            )
        else:
            chunk_ids = [row["id"] for row in connection.execute("SELECT id FROM chunks WHERE doc_id=?", (doc_id,))]
            connection.execute("UPDATE chunks SET enabled=?, update_time=? WHERE doc_id=?", (state, now_iso(), doc_id))
    if state:
        index_chunks(vector_rows(chunk_ids or []))
    else:
        remove_chunks(chunk_ids or [])
    return ok()


@app.post(PREFIX + "/knowledge-base/docs/{doc_id}/chunks/batch-enable")
def batch_enable(doc_id: str, payload: dict[str, Any] = Body(default={}), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return _set_chunk_state(doc_id, payload.get("chunkIds"), 1)


@app.post(PREFIX + "/knowledge-base/docs/{doc_id}/chunks/batch-disable")
def batch_disable(doc_id: str, payload: dict[str, Any] = Body(default={}), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return _set_chunk_state(doc_id, payload.get("chunkIds"), 0)


@app.get(PREFIX + "/knowledge-base/docs/{doc_id}/chunk-logs")
def chunk_logs(doc_id: str, current: int = 1, size: int = 10, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        total = connection.execute("SELECT COUNT(*) FROM chunk_logs WHERE doc_id=?", (doc_id,)).fetchone()[0]
        rows = connection.execute("SELECT * FROM chunk_logs WHERE doc_id=? ORDER BY create_time DESC LIMIT ? OFFSET ?", (doc_id, size, (current - 1) * size)).fetchall()
    records = []
    for raw in rows:
        row = dict(raw)
        records.append({
            "id": row["id"], "docId": row["doc_id"], "status": row["status"], "processMode": row["process_mode"],
            "chunkStrategy": row["chunk_strategy"], "pipelineId": row["pipeline_id"], "pipelineName": row["pipeline_name"],
            "extractDuration": row["extract_duration"], "chunkDuration": row["chunk_duration"], "embedDuration": row["embed_duration"],
            "persistDuration": row["persist_duration"], "otherDuration": row["other_duration"], "totalDuration": row["total_duration"],
            "chunkCount": row["chunk_count"], "errorMessage": row["error_message"], "startTime": row["start_time"],
            "endTime": row["end_time"], "createTime": row["create_time"],
        })
    return ok(page(records, total, current, size))


def sse(event: str, payload: Any) -> str:
    data = json.dumps(payload, ensure_ascii=False) if not isinstance(payload, str) else payload
    return f"event: {event}\ndata: {data}\n\n"


def persist_assistant_message(conversation_id: str, content: str) -> str:
    with db.connect() as connection:
        cursor = connection.execute(
            "INSERT INTO messages(conversation_id, role, content, create_time) VALUES (?, 'assistant', ?, ?)",
            (conversation_id, content, now_iso()),
        )
        connection.execute("UPDATE conversations SET update_time=? WHERE id=?", (now_iso(), conversation_id))
    return str(cursor.lastrowid)


def trace_id_for_task(task_id: str) -> str | None:
    with db.connect() as connection:
        row = connection.execute(
            "SELECT trace_id FROM traces WHERE task_id=? ORDER BY start_time DESC LIMIT 1",
            (task_id,),
        ).fetchone()
    return row["trace_id"] if row else None


@app.get(PREFIX + "/rag/v3/chat")
async def rag_chat(
    question: str,
    conversationId: str | None = None,
    deepThinking: bool = False,
    user: dict[str, Any] = Depends(current_user),
) -> StreamingResponse:
    question = question.strip()
    if not question:
        fail("问题不能为空")
    conversation_id, task_id = conversationId or new_id(), new_id()
    with db.connect() as connection:
        conversation = connection.execute("SELECT * FROM conversations WHERE id=?", (conversation_id,)).fetchone()
        if conversation and conversation["user_id"] != user["id"]:
            fail("无权访问该会话", 403)
        history_rows = connection.execute(
            "SELECT role, content FROM messages WHERE conversation_id=? ORDER BY id DESC LIMIT 8",
            (conversation_id,),
        ).fetchall()
        history = [dict(row) for row in reversed(history_rows)]
        timestamp = now_iso()
        if not conversation:
            title = question[:30]
            connection.execute("INSERT INTO conversations VALUES (?, ?, ?, ?, ?)", (conversation_id, user["id"], title, timestamp, timestamp))
        else:
            title = conversation["title"]
            connection.execute("UPDATE conversations SET update_time=? WHERE id=?", (timestamp, conversation_id))
        connection.execute("INSERT INTO messages(conversation_id, role, content, create_time) VALUES (?, 'user', ?, ?)", (conversation_id, question, timestamp))

    async def generate():
        control = GenerationControl(
            task_id=task_id,
            user_id=user["id"],
            conversation_id=conversation_id,
            loop=asyncio.get_running_loop(),
        )
        generation_tasks.register(control)
        token_queue: asyncio.Queue[str | None] = asyncio.Queue()
        emitted = ""
        agent_task: asyncio.Task[Any] | None = None

        async def forward_token(delta: str) -> None:
            if not control.cancelled:
                await token_queue.put(delta)

        async def execute_agent():
            try:
                return await agent.run(
                    question,
                    conversation_id,
                    task_id,
                    user,
                    history,
                    deepThinking,
                    on_token=forward_token,
                    cancelled=lambda: control.cancelled,
                )
            finally:
                await token_queue.put(None)

        try:
            agent_task = asyncio.create_task(execute_agent(), name=f"generation-{task_id}")
            generation_tasks.bind(task_id, agent_task)
            await asyncio.sleep(0)
            yield sse("meta", {"conversationId": conversation_id, "taskId": task_id})
            if deepThinking:
                yield sse("message", {"type": "think", "delta": "正在规划工具、筛选证据并组装上下文。"})
            while True:
                delta = await token_queue.get()
                if delta is None:
                    break
                if control.cancelled:
                    continue
                emitted += delta
                yield sse("message", {"type": "response", "delta": delta})
            result = await agent_task
            if control.cancelled:
                raise asyncio.CancelledError
            message_id = persist_assistant_message(conversation_id, emitted)
            payload = {
                "messageId": message_id, "title": title, "traceId": result.trace_id,
                "modelRoute": result.model_route, "sources": [hit.as_dict() for hit in result.hits],
            }
            yield sse("finish", payload)
            yield sse("done", {})
        except asyncio.CancelledError:
            if not control.cancelled:
                generation_tasks.disconnect(task_id)
                cancelled_content = emitted.rstrip()
                if cancelled_content:
                    persist_assistant_message(conversation_id, cancelled_content + "\n\n（连接中断，已停止生成）")
                raise
            cancelled_content = emitted.rstrip()
            stored_content = (
                cancelled_content + "\n\n（已停止生成）" if cancelled_content else "（已停止生成）"
            )
            message_id = persist_assistant_message(conversation_id, stored_content)
            yield sse("cancel", {
                "messageId": message_id,
                "title": title,
                "traceId": trace_id_for_task(task_id),
                "modelRoute": "cancelled",
                "sources": [],
                "reason": control.cancel_reason,
            })
            yield sse("done", {})
        except Exception as exc:
            yield sse("error", {"error": str(exc)})
        finally:
            if agent_task and not agent_task.done():
                generation_tasks.disconnect(task_id)
                agent_task.cancel()
                with suppress(asyncio.CancelledError):
                    await agent_task
            generation_tasks.unregister(task_id, control)

    return StreamingResponse(generate(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post(PREFIX + "/rag/v3/stop")
async def stop_task(taskId: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    task_id = taskId.strip()
    if not task_id:
        fail("任务 ID 不能为空")
    status = generation_tasks.cancel(task_id, user["id"])
    if status == "forbidden":
        fail("无权停止该生成任务", 403)
    return ok({"taskId": task_id, "status": status})


@app.get(PREFIX + "/conversations")
def conversations(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    with db.connect() as connection:
        rows = connection.execute("SELECT * FROM conversations WHERE user_id=? ORDER BY update_time DESC", (user["id"],)).fetchall()
    return ok([{"conversationId": row["id"], "title": row["title"], "lastTime": row["update_time"]} for row in rows])


@app.put(PREFIX + "/conversations/{conversation_id}")
def rename_conversation(conversation_id: str, payload: dict[str, Any] = Body(...), user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("UPDATE conversations SET title=?, update_time=? WHERE id=? AND user_id=?", (str(payload.get("title", "新对话"))[:60], now_iso(), conversation_id, user["id"]))
    return ok()


@app.delete(PREFIX + "/conversations/{conversation_id}")
def delete_conversation(conversation_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("DELETE FROM conversations WHERE id=? AND user_id=?", (conversation_id, user["id"]))
    return ok()


@app.get(PREFIX + "/conversations/{conversation_id}/messages")
def messages(conversation_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    with db.connect() as connection:
        allowed = connection.execute("SELECT 1 FROM conversations WHERE id=? AND user_id=?", (conversation_id, user["id"])).fetchone()
        if not allowed:
            fail("会话不存在", 404)
        rows = connection.execute("SELECT * FROM messages WHERE conversation_id=? ORDER BY id", (conversation_id,)).fetchall()
    return ok([{"id": row["id"], "conversationId": row["conversation_id"], "role": row["role"], "content": row["content"], "vote": row["vote"], "createTime": row["create_time"]} for row in rows])


@app.post(PREFIX + "/conversations/messages/{message_id}/feedback")
def feedback(message_id: int, payload: dict[str, Any] = Body(...), user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute(
            """UPDATE messages SET vote=? WHERE id=? AND conversation_id IN
               (SELECT id FROM conversations WHERE user_id=?)""",
            (payload.get("vote"), message_id, user["id"]),
        )
    return ok()


def sample_api(row: dict[str, Any]) -> dict[str, Any]:
    return {"id": row["id"], "title": row["title"], "description": row["description"], "question": row["question"], "createTime": row["create_time"], "updateTime": row["update_time"]}


@app.get(PREFIX + "/rag/sample-questions")
def public_samples(_: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    with db.connect() as connection:
        rows = connection.execute("SELECT * FROM sample_questions ORDER BY create_time").fetchall()
    return ok([sample_api(dict(row)) for row in rows])


@app.get(PREFIX + "/sample-questions")
def sample_page(current: int = 1, size: int = 10, keyword: str | None = None, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    where, values = (" WHERE question LIKE ? OR title LIKE ?", [f"%{keyword}%", f"%{keyword}%"]) if keyword else ("", [])
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM sample_questions{where}", values).fetchone()[0]
        rows = connection.execute(f"SELECT * FROM sample_questions{where} ORDER BY create_time DESC LIMIT ? OFFSET ?", (*values, size, (current - 1) * size)).fetchall()
    return ok(page([sample_api(dict(row)) for row in rows], total, current, size))


@app.get(PREFIX + "/sample-questions/{item_id}")
def sample_detail(item_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM sample_questions WHERE id=?", (item_id,)).fetchone()
    if not row:
        fail("示例问题不存在", 404)
    return ok(sample_api(dict(row)))


@app.post(PREFIX + "/sample-questions")
def create_sample(payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    item_id, timestamp = new_id(), now_iso()
    with db.connect() as connection:
        connection.execute("INSERT INTO sample_questions VALUES (?, ?, ?, ?, ?, ?)", (item_id, payload.get("title"), payload.get("description"), payload.get("question", ""), timestamp, timestamp))
    return ok(item_id)


@app.put(PREFIX + "/sample-questions/{item_id}")
def update_sample(item_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        existing = connection.execute("SELECT * FROM sample_questions WHERE id=?", (item_id,)).fetchone()
        if not existing:
            fail("示例问题不存在", 404)
        connection.execute(
            "UPDATE sample_questions SET title=?, description=?, question=?, update_time=? WHERE id=?",
            (payload.get("title", existing["title"]), payload.get("description", existing["description"]), payload.get("question", existing["question"]), now_iso(), item_id),
        )
    return ok()


@app.delete(PREFIX + "/sample-questions/{item_id}")
def delete_sample(item_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("DELETE FROM sample_questions WHERE id=?", (item_id,))
    return ok()


def mapping_api(row: dict[str, Any]) -> dict[str, Any]:
    return {"id": row["id"], "sourceTerm": row["source_term"], "targetTerm": row["target_term"], "matchType": row["match_type"], "priority": row["priority"], "enabled": bool(row["enabled"]), "remark": row["remark"], "createTime": row["create_time"], "updateTime": row["update_time"]}


@app.get(PREFIX + "/mappings")
def mappings(current: int = 1, size: int = 10, keyword: str | None = None, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    where, values = (" WHERE source_term LIKE ? OR target_term LIKE ?", [f"%{keyword}%", f"%{keyword}%"]) if keyword else ("", [])
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM mappings{where}", values).fetchone()[0]
        rows = connection.execute(f"SELECT * FROM mappings{where} ORDER BY priority DESC, create_time DESC LIMIT ? OFFSET ?", (*values, size, (current - 1) * size)).fetchall()
    return ok(page([mapping_api(dict(row)) for row in rows], total, current, size))


@app.get(PREFIX + "/mappings/{item_id}")
def mapping_detail(item_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM mappings WHERE id=?", (item_id,)).fetchone()
    if not row:
        fail("映射不存在", 404)
    return ok(mapping_api(dict(row)))


@app.post(PREFIX + "/mappings")
def create_mapping(payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    item_id, timestamp = new_id(), now_iso()
    with db.connect() as connection:
        connection.execute(
            "INSERT INTO mappings VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (item_id, payload.get("sourceTerm", ""), payload.get("targetTerm", ""), payload.get("matchType", 1), payload.get("priority", 0), int(payload.get("enabled", True)), payload.get("remark"), timestamp, timestamp),
        )
    return ok(item_id)


@app.put(PREFIX + "/mappings/{item_id}")
def update_mapping(item_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM mappings WHERE id=?", (item_id,)).fetchone()
        if not row:
            fail("映射不存在", 404)
        connection.execute(
            """UPDATE mappings SET source_term=?, target_term=?, match_type=?, priority=?, enabled=?, remark=?, update_time=? WHERE id=?""",
            (payload.get("sourceTerm", row["source_term"]), payload.get("targetTerm", row["target_term"]), payload.get("matchType", row["match_type"]), payload.get("priority", row["priority"]), int(payload.get("enabled", bool(row["enabled"]))), payload.get("remark", row["remark"]), now_iso(), item_id),
        )
    return ok()


@app.delete(PREFIX + "/mappings/{item_id}")
def delete_mapping(item_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("DELETE FROM mappings WHERE id=?", (item_id,))
    return ok()


def intent_api(row: dict[str, Any], children: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    return {
        "id": row["id"], "intentCode": row["intent_code"], "name": row["name"], "level": row["level"],
        "parentCode": row["parent_code"], "description": row["description"], "examples": row["examples"],
        "collectionName": row["collection_name"], "mcpToolId": row["mcp_tool_id"], "topK": row["top_k"],
        "kind": row["kind"], "sortOrder": row["sort_order"], "enabled": row["enabled"],
        "promptSnippet": row["prompt_snippet"], "promptTemplate": row["prompt_template"],
        "paramPromptTemplate": row["param_prompt_template"], "children": children or [],
    }


@app.get(PREFIX + "/intent-tree/trees")
def intent_tree(_: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        rows = [dict(row) for row in connection.execute("SELECT * FROM intents ORDER BY sort_order, id").fetchall()]
    by_parent: dict[str | None, list[dict[str, Any]]] = {}
    for row in rows:
        by_parent.setdefault(row["parent_code"], []).append(row)
    def build(row: dict[str, Any]) -> dict[str, Any]:
        return intent_api(row, [build(child) for child in by_parent.get(row["intent_code"], [])])
    return ok([build(row) for row in by_parent.get(None, [])])


@app.post(PREFIX + "/intent-tree")
def create_intent(payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    examples = payload.get("examples")
    if isinstance(examples, list):
        examples = "\n".join(examples)
    with db.connect() as connection:
        cursor = connection.execute(
            """INSERT INTO intents(kb_id,intent_code,name,level,parent_code,description,examples,mcp_tool_id,top_k,kind,sort_order,enabled,prompt_snippet,prompt_template,param_prompt_template)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (payload.get("kbId"), payload.get("intentCode"), payload.get("name"), payload.get("level", 1), payload.get("parentCode"), payload.get("description"), examples, payload.get("mcpToolId"), payload.get("topK", 5), payload.get("kind", 0), payload.get("sortOrder", 0), payload.get("enabled", 1), payload.get("promptSnippet"), payload.get("promptTemplate"), payload.get("paramPromptTemplate")),
        )
    return ok(str(cursor.lastrowid))


@app.put(PREFIX + "/intent-tree/{item_id}")
def update_intent(item_id: int, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    mapping = {"name":"name","level":"level","parentCode":"parent_code","description":"description","collectionName":"collection_name","mcpToolId":"mcp_tool_id","topK":"top_k","kind":"kind","sortOrder":"sort_order","enabled":"enabled","promptSnippet":"prompt_snippet","promptTemplate":"prompt_template","paramPromptTemplate":"param_prompt_template"}
    fields, values = [], []
    for key, column in mapping.items():
        if key in payload:
            fields.append(f"{column}=?")
            values.append(payload[key])
    if "examples" in payload:
        fields.append("examples=?")
        values.append("\n".join(payload["examples"]) if isinstance(payload["examples"], list) else payload["examples"])
    if fields:
        values.append(item_id)
        with db.connect() as connection:
            connection.execute(f"UPDATE intents SET {', '.join(fields)} WHERE id=?", values)
    return ok()


@app.delete(PREFIX + "/intent-tree/{item_id}")
def delete_intent(item_id: int, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("DELETE FROM intents WHERE id=?", (item_id,))
    return ok()


def intent_batch(payload: dict[str, Any], action: str) -> dict[str, Any]:
    ids = payload.get("ids", [])
    if not ids:
        return ok()
    placeholders = ",".join("?" for _ in ids)
    with db.connect() as connection:
        if action == "delete":
            connection.execute(f"DELETE FROM intents WHERE id IN ({placeholders})", ids)
        else:
            connection.execute(f"UPDATE intents SET enabled=? WHERE id IN ({placeholders})", (1 if action == "enable" else 0, *ids))
    return ok()


@app.post(PREFIX + "/intent-tree/batch/enable")
def batch_enable_intents(payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return intent_batch(payload, "enable")


@app.post(PREFIX + "/intent-tree/batch/disable")
def batch_disable_intents(payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return intent_batch(payload, "disable")


@app.post(PREFIX + "/intent-tree/batch/delete")
def batch_delete_intents(payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    return intent_batch(payload, "delete")


def trace_api(row: dict[str, Any]) -> dict[str, Any]:
    return {"traceId": row["trace_id"], "traceName": row["trace_name"], "entryMethod": row["entry_method"], "conversationId": row["conversation_id"], "taskId": row["task_id"], "userName": row["username"], "username": row["username"], "userId": row["user_id"], "status": row["status"], "errorMessage": row["error_message"], "durationMs": row["duration_ms"], "startTime": row["start_time"], "endTime": row["end_time"]}


def trace_node_api(row: dict[str, Any]) -> dict[str, Any]:
    return {"traceId": row["trace_id"], "nodeId": row["node_id"], "parentNodeId": row["parent_node_id"], "depth": row["depth"], "nodeType": row["node_type"], "nodeName": row["node_name"], "className": row["class_name"], "methodName": row["method_name"], "status": row["status"], "errorMessage": row["error_message"], "durationMs": row["duration_ms"], "startTime": row["start_time"], "endTime": row["end_time"]}


@app.get(PREFIX + "/rag/traces/runs")
def trace_runs(current: int = 1, size: int = 10, traceId: str | None = None, conversationId: str | None = None, taskId: str | None = None, status: str | None = None, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    clauses, values = [], []
    for value, column in ((traceId,"trace_id"),(conversationId,"conversation_id"),(taskId,"task_id"),(status,"status")):
        if value:
            clauses.append(f"{column}=?")
            values.append(value)
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM traces{where}", values).fetchone()[0]
        rows = connection.execute(f"SELECT * FROM traces{where} ORDER BY start_time DESC LIMIT ? OFFSET ?", (*values, size, (current - 1) * size)).fetchall()
    return ok(page([trace_api(dict(row)) for row in rows], total, current, size))


@app.get(PREFIX + "/rag/traces/runs/{trace_id}")
def trace_detail(trace_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        run = connection.execute("SELECT * FROM traces WHERE trace_id=?", (trace_id,)).fetchone()
        nodes = connection.execute("SELECT * FROM trace_nodes WHERE trace_id=? ORDER BY id", (trace_id,)).fetchall()
    if not run:
        fail("Trace 不存在", 404)
    return ok({"run": trace_api(dict(run)), "nodes": [trace_node_api(dict(row)) for row in nodes]})


@app.get(PREFIX + "/rag/traces/runs/{trace_id}/nodes")
def trace_nodes(trace_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        rows = connection.execute("SELECT * FROM trace_nodes WHERE trace_id=? ORDER BY id", (trace_id,)).fetchall()
    return ok([trace_node_api(dict(row)) for row in rows])


@app.get(PREFIX + "/admin/dashboard/overview")
def dashboard_overview(window: str = "24h", _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    since = (datetime.now(UTC) - timedelta(hours=24)).isoformat(timespec="seconds")
    with db.connect() as connection:
        total_users = connection.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        sessions = connection.execute("SELECT COUNT(*) FROM conversations").fetchone()[0]
        sessions_24 = connection.execute("SELECT COUNT(*) FROM conversations WHERE create_time>=?", (since,)).fetchone()[0]
        messages_count = connection.execute("SELECT COUNT(*) FROM messages").fetchone()[0]
        messages_24 = connection.execute("SELECT COUNT(*) FROM messages WHERE create_time>=?", (since,)).fetchone()[0]
    kpi = lambda value: {"value": value, "delta": 0, "deltaPct": 0}
    return ok({"window": window, "compareWindow": "previous", "updatedAt": int(time.time() * 1000), "kpis": {"totalUsers": kpi(total_users), "activeUsers": kpi(total_users), "totalSessions": kpi(sessions), "sessions24h": kpi(sessions_24), "totalMessages": kpi(messages_count), "messages24h": kpi(messages_24)}})


@app.get(PREFIX + "/admin/dashboard/performance")
def dashboard_performance(window: str = "24h", _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        durations = [row[0] or 0 for row in connection.execute("SELECT duration_ms FROM traces WHERE status='success'").fetchall()]
        total = connection.execute("SELECT COUNT(*) FROM traces").fetchone()[0]
        errors = connection.execute("SELECT COUNT(*) FROM traces WHERE status='error'").fetchone()[0]
    durations.sort()
    avg = round(sum(durations) / len(durations)) if durations else 0
    p95 = durations[min(len(durations) - 1, int(len(durations) * 0.95))] if durations else 0
    error_rate = (errors / total * 100) if total else 0
    return ok({"window": window, "avgLatencyMs": avg, "p95LatencyMs": p95, "successRate": 100-error_rate, "errorRate": error_rate, "noDocRate": 0, "slowRate": 0})


@app.get(PREFIX + "/admin/dashboard/trends")
def dashboard_trends(metric: str, window: str = "7d", granularity: str = "day", _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        rows = connection.execute("SELECT substr(start_time,1,10) day, COUNT(*) count FROM traces GROUP BY day ORDER BY day").fetchall()
    points = [{"ts": int(datetime.fromisoformat(row["day"]).replace(tzinfo=UTC).timestamp() * 1000), "value": row["count"]} for row in rows if row["day"]]
    return ok({"metric": metric, "window": window, "granularity": granularity, "series": [{"name": metric, "data": points}]})


@app.get(PREFIX + "/rag/settings")
def rag_settings(_: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    configured = bool(settings.llm_api_key)
    embedding_configured = bool(settings.embedding_api_key)
    rerank_configured = bool(settings.rerank_api_key)
    candidates = [{"id": "local-grounded-fallback", "provider": "local", "model": "extractive-grounded", "priority": 100, "enabled": True, "supportsThinking": True}]
    if configured:
        candidates.insert(0, {"id": settings.llm_model, "provider": "openai-compatible", "model": settings.llm_model, "url": settings.llm_base_url, "priority": 0, "enabled": True, "supportsThinking": True})
    return ok({
        "upload": {"maxFileSize": 50 * 1024 * 1024, "maxRequestSize": 100 * 1024 * 1024},
        "rag": {"default": {"collectionName": settings.vector_collection, "dimension": settings.embedding_dimension, "metricType": "BM25_QDRANT_RRF_RERANK"}, "queryRewrite": {"enabled": settings.query_rewrite_enabled, "maxHistoryMessages": 8, "maxHistoryChars": 4800}, "rateLimit": {"global": {"enabled": False, "maxConcurrent": 20, "maxWaitSeconds": 3, "leaseSeconds": 60, "pollIntervalMs": 100}}, "memory": {"historyKeepTurns": 4, "summaryStartTurns": 8, "summaryEnabled": True, "ttlMinutes": 0, "summaryMaxChars": 600, "titleMaxLength": 30}},
        "ai": {"providers": {"openai-compatible": {"url": settings.llm_base_url, "apiKey": "已配置" if configured else None, "endpoints": {"chat": "/chat/completions"}}, "local": {"url": "in-process", "endpoints": {"chat": "grounded-fallback"}}}, "selection": {"failureThreshold": 1, "openDurationMs": 30000}, "stream": {"mode": "upstream-token-delta", "contextTokenBudget": settings.context_token_budget, "maxOutputTokens": settings.max_output_tokens}, "chat": {"defaultModel": settings.llm_model if configured else "local-grounded-fallback", "deepThinkingModel": settings.llm_model if configured else "local-grounded-fallback", "candidates": candidates}, "embedding": {"defaultModel": settings.embedding_model if embedding_configured else "local-hash-embedding", "candidates": [{"id": settings.embedding_model if embedding_configured else "local-hash-embedding", "provider": "openai-compatible" if embedding_configured else "local", "model": settings.embedding_model if embedding_configured else "deterministic-hash", "dimension": settings.embedding_dimension, "enabled": True}]}, "rerank": {"defaultModel": settings.rerank_model if rerank_configured else "local-feature-reranker", "candidates": [{"id": settings.rerank_model if rerank_configured else "local-feature-reranker", "provider": "model-api" if rerank_configured else "local", "model": settings.rerank_model if rerank_configured else "feature-reranker", "enabled": True}]}}
    })


def pipeline_api(row: dict[str, Any]) -> dict[str, Any]:
    nodes = loads(row.get("nodes_json"), [])
    for index, node in enumerate(nodes):
        node.setdefault("id", index + 1)
    return {"id": row["id"], "name": row["name"], "description": row["description"], "createdBy": row["created_by"], "nodes": nodes, "createTime": row["create_time"], "updateTime": row["update_time"]}


@app.get(PREFIX + "/ingestion/pipelines")
def pipelines(pageNo: int = 1, pageSize: int = 10, keyword: str | None = None, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    where, values = (" WHERE name LIKE ?", [f"%{keyword}%"]) if keyword else ("", [])
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM pipelines{where}", values).fetchone()[0]
        rows = connection.execute(f"SELECT * FROM pipelines{where} ORDER BY create_time DESC LIMIT ? OFFSET ?", (*values, pageSize, (pageNo - 1) * pageSize)).fetchall()
    return ok(page([pipeline_api(dict(row)) for row in rows], total, pageNo, pageSize))


@app.get(PREFIX + "/ingestion/pipelines/{pipeline_id}")
def pipeline_detail(pipeline_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM pipelines WHERE id=?", (pipeline_id,)).fetchone()
    if not row:
        fail("流水线不存在", 404)
    return ok(pipeline_api(dict(row)))


@app.post(PREFIX + "/ingestion/pipelines")
def create_pipeline(payload: dict[str, Any] = Body(...), user: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    item_id, timestamp = new_id(), now_iso()
    with db.connect() as connection:
        connection.execute("INSERT INTO pipelines VALUES (?, ?, ?, ?, ?, ?, ?)", (item_id, payload.get("name", "未命名流水线"), payload.get("description"), user["username"], dumps(payload.get("nodes", [])), timestamp, timestamp))
        row = connection.execute("SELECT * FROM pipelines WHERE id=?", (item_id,)).fetchone()
    return ok(pipeline_api(dict(row)))


@app.put(PREFIX + "/ingestion/pipelines/{pipeline_id}")
def update_pipeline(pipeline_id: str, payload: dict[str, Any] = Body(...), _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("UPDATE pipelines SET name=?,description=?,nodes_json=?,update_time=? WHERE id=?", (payload.get("name", "未命名流水线"), payload.get("description"), dumps(payload.get("nodes", [])), now_iso(), pipeline_id))
        row = connection.execute("SELECT * FROM pipelines WHERE id=?", (pipeline_id,)).fetchone()
    return ok(pipeline_api(dict(row)))


@app.delete(PREFIX + "/ingestion/pipelines/{pipeline_id}")
def delete_pipeline(pipeline_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        connection.execute("DELETE FROM pipelines WHERE id=?", (pipeline_id,))
    return ok()


def task_api(row: dict[str, Any]) -> dict[str, Any]:
    return {"id": row["id"], "pipelineId": row["pipeline_id"], "sourceType": row["source_type"], "sourceLocation": row["source_location"], "sourceFileName": row["source_file_name"], "status": row["status"], "chunkCount": row["chunk_count"], "errorMessage": row["error_message"], "logs": loads(row["logs_json"], []), "metadata": loads(row["metadata_json"], {}), "startedAt": row["started_at"], "completedAt": row["completed_at"], "createdBy": row["created_by"], "createTime": row["create_time"], "updateTime": row["update_time"]}


def store_ingestion_task(pipeline_id: str, source_type: str, location: str, filename: str | None, text: str, user: dict[str, Any]) -> dict[str, Any]:
    item_id, timestamp = new_id(), now_iso()
    chunks_count = len(structure_chunks(text))
    logs = [
        {"nodeId":"parse","nodeType":"parser","message":"文档解析完成","durationMs":1,"success":True},
        {"nodeId":"chunk","nodeType":"chunker","message":f"生成 {chunks_count} 个分块","durationMs":1,"success":True},
        {"nodeId":"index","nodeType":"indexer","message":"等待绑定知识库后建立索引","durationMs":0,"success":True},
    ]
    with db.connect() as connection:
        connection.execute("INSERT INTO ingestion_tasks VALUES (?, ?, ?, ?, ?, 'completed', ?, NULL, ?, ?, ?, ?, ?, ?, ?)", (item_id, pipeline_id, source_type, location, filename, chunks_count, dumps(logs), dumps({}), timestamp, now_iso(), user["username"], timestamp, now_iso()))
    return {"taskId": item_id, "pipelineId": pipeline_id, "status": "completed", "chunkCount": chunks_count, "message": "入库流水线执行完成"}


@app.post(PREFIX + "/ingestion/tasks")
def create_ingestion_task(payload: dict[str, Any] = Body(...), user: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    source = payload.get("source", {})
    location = str(source.get("location", ""))
    text = f"Source: {location}\nMetadata: {json.dumps(payload.get('metadata', {}), ensure_ascii=False)}"
    return ok(store_ingestion_task(str(payload.get("pipelineId")), str(source.get("type", "text")), location, source.get("fileName"), text, user))


@app.post(PREFIX + "/ingestion/tasks/upload")
async def upload_ingestion_task(pipelineId: str, file: UploadFile = File(...), user: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    text, raw = await upload_to_text(file)
    path = settings.upload_dir / f"{new_id()}_{Path(file.filename or 'document').name}"
    path.write_bytes(raw)
    return ok(store_ingestion_task(pipelineId, "file", str(path), file.filename, text, user))


@app.get(PREFIX + "/ingestion/tasks")
def ingestion_tasks(pageNo: int = 1, pageSize: int = 10, status: str | None = None, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    where, values = (" WHERE status=?", [status]) if status else ("", [])
    with db.connect() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM ingestion_tasks{where}", values).fetchone()[0]
        rows = connection.execute(f"SELECT * FROM ingestion_tasks{where} ORDER BY create_time DESC LIMIT ? OFFSET ?", (*values, pageSize, (pageNo - 1) * pageSize)).fetchall()
    return ok(page([task_api(dict(row)) for row in rows], total, pageNo, pageSize))


@app.get(PREFIX + "/ingestion/tasks/{task_id}")
def ingestion_task(task_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM ingestion_tasks WHERE id=?", (task_id,)).fetchone()
    if not row:
        fail("入库任务不存在", 404)
    return ok(task_api(dict(row)))


@app.get(PREFIX + "/ingestion/tasks/{task_id}/nodes")
def ingestion_task_nodes(task_id: str, _: dict[str, Any] = Depends(admin_user)) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT * FROM ingestion_tasks WHERE id=?", (task_id,)).fetchone()
    if not row:
        fail("入库任务不存在", 404)
    logs = loads(row["logs_json"], [])
    return ok([{"id": f"{task_id}-{index}", "taskId": task_id, "pipelineId": row["pipeline_id"], "nodeId": log.get("nodeId"), "nodeType": log.get("nodeType"), "nodeOrder": index, "status": "success" if log.get("success") else "error", "durationMs": log.get("durationMs"), "message": log.get("message"), "errorMessage": log.get("error"), "output": {}, "createTime": row["create_time"], "updateTime": row["update_time"]} for index, log in enumerate(logs)])
